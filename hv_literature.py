#!/usr/bin/env python3
"""Local evidence library and bounded OpenAlex literature watch. No LLM keys required."""
from __future__ import annotations

import argparse
import collections
import contextlib
import datetime as dt
import fcntl
import hashlib
import html
import json
import logging
import os
from pathlib import Path
import re
import sqlite3
import sys
import time
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
from zoneinfo import ZoneInfo

ROOT = Path(__file__).resolve().parent
VERSION = "0.2.0"
ROLE_CUES = {
    "advantages": ["advantage", "efficient", "efficiency", "robust", "benefit", "outperform", "superior", "accurate", "优势", "优点", "高效", "鲁棒", "准确"],
    "limitations": ["limitation", "limited", "drawback", "however", "challenge", "difficult", "suffer", "expensive", "fails", "disadvantage", "缺点", "局限", "不足", "难以", "然而", "耗时"],
    "methods": ["propose", "framework", "method", "model", "formulation", "approach", "方法", "模型", "提出", "构建"],
    "results": ["results", "demonstrate", "experiment", "simulation", "evaluation", "结果", "试验", "仿真", "验证"],
    "background": ["recent", "increasing", "important", "motivation", "背景", "近年来", "重要"]
}
ROLE_ZH = {"advantages":"优势", "limitations":"局限", "methods":"方法", "results":"实验结果", "background":"研究背景", "all":"不限"}
CONCEPTS = {
    "全纯": ["holomorphic", "HEM", "HELM", "全纯"],
    "潮流": ["power flow", "load flow", "潮流"],
    "收敛": ["convergence", "converge", "收敛"],
    "稳定": ["stability", "stable", "稳定"],
    "计算复杂度": ["computational complexity", "computational burden", "计算复杂度"],
    "绝缘气体": ["insulating gas", "dielectric gas", "绝缘气体"],
    "分子": ["molecular", "molecule", "分子"],
    "构效": ["structure property", "QSAR", "QSPR", "构效"],
    "小样本": ["small sample", "few shot", "limited data", "小样本"],
    "局部放电": ["partial discharge", "局部放电"],
    "多模态": ["multimodal", "multi-modal", "multi sensor", "多模态"],
    "声光电": ["acoustic", "optical", "electrical", "声光电"],
    "融合": ["fusion", "融合"],
    "故障": ["fault", "defect", "故障", "缺陷"],
    "预测": ["prediction", "forecast", "prognostic", "预测"],
}


def now():
    return dt.datetime.now(ZoneInfo("Asia/Shanghai"))


def dump(obj):
    return json.dumps(obj, ensure_ascii=False, indent=2)


def normalize(s):
    s = unicodedata.normalize("NFKC", s or "").replace("\x00", "")
    s = re.sub(r"(?<=[A-Za-z])-\s*\n\s*(?=[a-z])", "", s)
    return re.sub(r"\s+", " ", s).strip()


def index_text(s):
    # English porter stemming + explicit Chinese bigrams in the same FTS index.
    s = normalize(s).lower()
    grams = [v[i:i+2] for v in re.findall(r"[\u4e00-\u9fff]+", s) for i in range(max(0, len(v)-1))]
    return s + " " + " ".join(grams)


def norm_journal(s):
    return re.sub(r"[^a-z0-9\u4e00-\u9fff]", "", s.lower().replace("&", "and"))


def atomic_write(path, text):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + f".{os.getpid()}.tmp")
    tmp.write_text(text, encoding="utf-8")
    tmp.replace(path)


class Library:
    def __init__(self, root=ROOT, config=None):
        self.root = Path(root).expanduser().resolve()
        if config is None:
            path = self.root / "config.json"
            if not path.is_file():
                raise RuntimeError("尚未初始化，请先运行 ./hv-literature init --library-root /你的论文目录")
            config = json.loads(path.read_text(encoding="utf-8"))
        self.config = config
        self.data = self.root / "data"
        self.data.mkdir(parents=True, exist_ok=True)
        self.db_path = self.data / "library.sqlite3"
        self.db = sqlite3.connect(self.db_path, timeout=30)
        self.db.row_factory = sqlite3.Row
        self.db.execute("PRAGMA journal_mode=WAL")
        self.db.execute("PRAGMA foreign_keys=ON")
        self.db.executescript("""
          CREATE TABLE IF NOT EXISTS documents(
            id TEXT PRIMARY KEY, title TEXT, doi TEXT, pages INTEGER,
            low_text_pages TEXT, indexed_at TEXT);
          CREATE TABLE IF NOT EXISTS files(
            path TEXT PRIMARY KEY, document_id TEXT REFERENCES documents(id),
            mtime_ns INTEGER, size INTEGER, active INTEGER DEFAULT 1);
          CREATE TABLE IF NOT EXISTS pages(
            document_id TEXT REFERENCES documents(id), page INTEGER, label TEXT, text TEXT,
            PRIMARY KEY(document_id,page));
          CREATE TABLE IF NOT EXISTS chunks(
            id INTEGER PRIMARY KEY, document_id TEXT REFERENCES documents(id),
            page INTEGER, ordinal INTEGER, text TEXT, roles TEXT);
          CREATE VIRTUAL TABLE IF NOT EXISTS chunks_fts USING fts5(title,body,tokenize='porter unicode61');
          CREATE TABLE IF NOT EXISTS journals(
            name TEXT, norm_name TEXT, issns TEXT, field TEXT, tier TEXT,
            source_path TEXT, source_sheet TEXT, source_row INTEGER,
            PRIMARY KEY(source_path,source_sheet,source_row));
          CREATE TABLE IF NOT EXISTS works(
            id TEXT PRIMARY KEY, canonical TEXT UNIQUE, title TEXT, publication_date TEXT,
            first_seen TEXT, metadata TEXT);
          CREATE TABLE IF NOT EXISTS featured(
            work_id TEXT REFERENCES works(id), date TEXT, topic TEXT,
            PRIMARY KEY(work_id,date));
          CREATE TABLE IF NOT EXISTS daily_runs(
            date TEXT PRIMARY KEY, status TEXT, report TEXT, payload TEXT);
          CREATE TABLE IF NOT EXISTS events(
            ts TEXT, kind TEXT, message TEXT);
        """)

    @contextlib.contextmanager
    def lock(self, name):
        with (self.data / f"{name}.lock").open("w") as f:
            try:
                fcntl.flock(f, fcntl.LOCK_EX | fcntl.LOCK_NB)
            except BlockingIOError:
                raise RuntimeError("已有同类任务运行，请稍后重试") from None
            yield

    def import_journals(self):
        if not self.config.get("journal_file"):
            return {"skipped":True,"reason":"未配置期刊目录；不影响论文检索和日报"}
        from openpyxl import load_workbook
        path = Path(self.config["journal_file"]).expanduser()
        wb = load_workbook(path, read_only=True, data_only=True)
        sh = wb[self.config.get("journal_sheet", "期刊目录")]
        header = [c.value for c in next(sh.iter_rows())]
        if header[1] != "刊名" or header[3] != "ISSN/EISSN" or header[6] != "期刊层级":
            raise ValueError("期刊表列结构改变，尚未导入")
        rows = []
        for n, cells in enumerate(sh.iter_rows(min_row=2, values_only=True), 2):
            if not cells[1]:
                continue
            title = str(cells[1]).strip()
            issns = re.findall(r"\d{4}-[\dXx]{4}", str(cells[3] or ""))
            rows.append((title, norm_journal(title), dump(issns), str(cells[5] or ""), str(cells[6] or ""), str(path), sh.title, n))
        wb.close()
        with self.db:
            self.db.execute("DELETE FROM journals")
            self.db.executemany("INSERT INTO journals VALUES (?,?,?,?,?,?,?,?)", rows)
        return {"journals":len(rows), "source":str(path), "sheet":sh.title}

    def ingest(self):
        from pypdf import PdfReader
        logging.getLogger("pypdf").setLevel(logging.ERROR)
        with self.lock("ingest"):
            stats = collections.Counter()
            errors, warnings = [], []
            seen = set()
            roots = [Path(x).expanduser().resolve() for x in self.config["library_roots"]]
            for base in roots:
                if not base.is_dir():
                    warnings.append(f"目录不可读，本轮保留旧索引：{base}")
                    continue
                paths = sorted(p for p in base.rglob("*") if p.is_file() and p.suffix.lower() == ".pdf")
                for path in paths:
                    path = path.resolve()
                    if not path.is_relative_to(base) or str(path) in seen:
                        continue
                    seen.add(str(path)); stats["pdf_files"] += 1
                    st = path.stat()
                    old = self.db.execute("SELECT * FROM files WHERE path=?", (str(path),)).fetchone()
                    if old and old["mtime_ns"] == st.st_mtime_ns and old["size"] == st.st_size:
                        self.db.execute("UPDATE files SET active=1 WHERE path=?", (str(path),)); self.db.commit()
                        stats["unchanged"] += 1
                        continue
                    sha = hashlib.sha256(path.read_bytes()).hexdigest()
                    try:
                        with self.db:
                            if not self.db.execute("SELECT 1 FROM documents WHERE id=?", (sha,)).fetchone():
                                reader = PdfReader(path, strict=False)
                                if reader.is_encrypted and not reader.decrypt(""):
                                    raise ValueError("encrypted_pdf")
                                page_text = [normalize(p.extract_text() or "") for p in reader.pages]
                                title = normalize((reader.metadata or {}).get("/Title", ""))
                                if len(title) < 10 or len(title) > 350 or title.lower().startswith(("microsoft", "untitled")):
                                    title = path.stem
                                doi_match = re.search(r"10\.\d{4,9}/[-._;()/:A-Z0-9]+", " ".join(page_text[:2]), re.I)
                                doi = doi_match.group().rstrip(".,;)") if doi_match else ""
                                low = [i for i, t in enumerate(page_text, 1) if len(t) < 50]
                                self.db.execute("INSERT INTO documents VALUES (?,?,?,?,?,?)", (sha,title,doi,len(page_text),dump(low),now().isoformat()))
                                try: labels = reader.page_labels
                                except Exception: labels = [str(i+1) for i in range(len(page_text))]
                                for n, txt in enumerate(page_text, 1):
                                    self.db.execute("INSERT INTO pages VALUES (?,?,?,?)", (sha,n,labels[n-1],txt))
                                    for k, chunk in enumerate(split_chunks(txt)):
                                        roles = [role for role,cues in ROLE_CUES.items() if any(cue in chunk.lower() for cue in cues)]
                                        cur = self.db.execute("INSERT INTO chunks(document_id,page,ordinal,text,roles) VALUES(?,?,?,?,?)", (sha,n,k,chunk,dump(roles)))
                                        self.db.execute("INSERT INTO chunks_fts(rowid,title,body) VALUES(?,?,?)", (cur.lastrowid,index_text(title),index_text(chunk)))
                                stats["new_documents"] += 1
                                if low: warnings.append(f"{path.name}：{len(low)} 页文字稀少，需人工核查或 OCR")
                            else:
                                stats["duplicate_files"] += 1
                            self.db.execute("INSERT OR REPLACE INTO files VALUES (?,?,?,?,1)", (str(path),sha,st.st_mtime_ns,st.st_size))
                    except Exception as e:
                        # A changed file that fails parsing must not retain stale searchable content.
                        with self.db: self.db.execute("UPDATE files SET active=0 WHERE path=?", (str(path),))
                        errors.append({"path":str(path),"type":type(e).__name__})
                with self.db:
                    for row in self.db.execute("SELECT path FROM files WHERE active=1").fetchall():
                        p = Path(row[0])
                        if p.is_relative_to(base) and row[0] not in seen:
                            self.db.execute("UPDATE files SET active=0 WHERE path=?",(row[0],))
            result = {"counts":dict(stats),"errors":errors,"warnings":warnings,"finished_at":now().isoformat()}
            atomic_write(self.data / "last_ingest.json", dump(result))
            return result

    def status(self):
        counts = {}
        for table in ["documents","files","pages","chunks","journals","works","featured"]:
            counts[table] = self.db.execute(f"SELECT count(*) FROM {table}").fetchone()[0]
        counts["active_files"] = self.db.execute("SELECT count(*) FROM files WHERE active=1").fetchone()[0]
        counts["active_documents"] = self.db.execute("SELECT count(DISTINCT document_id) FROM files WHERE active=1").fetchone()[0]
        counts["low_text_pages"] = sum(len(json.loads(r[0])) for r in self.db.execute("SELECT low_text_pages FROM documents"))
        return {"version":VERSION,"database":str(self.db_path),"counts":counts,"roots":self.config["library_roots"],
                "last_ingest":json.loads((self.data/"last_ingest.json").read_text()) if (self.data/"last_ingest.json").exists() else None}

    def search(self, query, role="all", limit=8):
        if role not in ROLE_ZH: raise ValueError("未知段落意图")
        limit = min(max(int(limit),1),30)
        query = str(query)[:600].strip()
        terms, groups = query_terms(query)
        if not terms: return {"query":query,"results":[],"note":"请提供技术关键词。"}
        match = " OR ".join('"'+t.replace('"','""')+'"' for t in terms)
        rows = self.db.execute("""SELECT c.*,d.title,d.doi,bm25(chunks_fts,2,1) AS rank
             FROM chunks_fts JOIN chunks c ON c.id=chunks_fts.rowid JOIN documents d ON d.id=c.document_id
             WHERE chunks_fts MATCH ? AND EXISTS(SELECT 1 FROM files f WHERE f.document_id=d.id AND active=1)
             ORDER BY rank LIMIT 350""", (match,)).fetchall()
        ranked=[]
        for row in rows:
            d=dict(row); text=d["text"].lower(); content=(d["title"]+" "+text).lower()
            if groups and not all(any(t.lower() in content for t in group) for group in groups): continue
            roles=json.loads(d["roles"])
            if role != "all" and role not in roles: continue
            d["score"] = -d.pop("rank") + 2*sum(t.lower() in text for t in terms)
            ranked.append(d)
        ranked.sort(key=lambda d:d["score"],reverse=True)
        selected=[]; used=collections.Counter(); seen_quotes=set()
        for d in ranked:
            if used[d["document_id"]]>=2:continue
            quote=best_sentence(d["text"],role,terms)
            if quote in seen_quotes: continue
            paths=[r[0] for r in self.db.execute("SELECT path FROM files WHERE document_id=? AND active=1 ORDER BY path",(d["document_id"],))]
            if not paths:continue
            d.update(paths=paths,quote=quote,roles=json.loads(d["roles"]),pdf_url=Path(paths[0]).as_uri()+f'#page={d["page"]}',
                     citation=f'{d["title"]}，PDF 第 {d["page"]} 页',basis="PDF文字层提取；可能存在分栏或公式识别误差")
            selected.append(d);seen_quotes.add(quote);used[d["document_id"]]+=1
            if len(selected)>=limit:break
        return {"query":query,"role":role,"expanded_terms":terms,"results":selected,
                "note":"原句是文字层摘录；roles 为关键词线索。写作组织与改写应由 Agent 结合整页核验，不能把模板当论文结论。"}

    def page(self, document_id, page):
        row=self.db.execute("SELECT * FROM pages WHERE document_id=? AND page=?",(document_id,int(page))).fetchone()
        if not row:raise ValueError("文档或页码不存在")
        result=dict(row)
        result["paths"]=[r[0] for r in self.db.execute("SELECT path FROM files WHERE document_id=? AND active=1",(document_id,))]
        return result

    def journal_match(self, work):
        source=((work.get("primary_location") or {}).get("source") or {})
        issns=set(source.get("issn") or [])
        name=norm_journal(source.get("display_name") or "")
        for row in self.db.execute("SELECT * FROM journals"):
            if issns.intersection(json.loads(row["issns"])) or (name and name==row["norm_name"]): return dict(row)
        return None

    def api(self, params):
        key=os.environ.get("OPENALEX_API_KEY")
        if not key:
            path=Path(self.config["credential_file"]).expanduser()
            if path.exists():key=json.loads(path.read_text()).get("openalex_api_key")
        if not key: raise RuntimeError("尚未配置 OpenAlex 凭据")
        url="https://api.openalex.org/works?"+urllib.parse.urlencode(params)
        request=urllib.request.Request(url,headers={"Authorization":"Bearer "+key,"User-Agent":"HVResearch/0.1"})
        for attempt in range(3):
            try:
                with urllib.request.urlopen(request,timeout=25) as r:return json.load(r)
            except urllib.error.HTTPError as e:
                if e.code in (429,500,502,503,504) and attempt<2:
                    time.sleep(2**attempt);continue
                raise RuntimeError(f"OpenAlex HTTP {e.code}（本轮不自动购买额度）") from None
            except (TimeoutError,urllib.error.URLError):
                if attempt<2:time.sleep(2**attempt);continue
                raise RuntimeError("OpenAlex 网络超时或连接失败") from None

    def daily(self, report_date=None):
        day=dt.date.fromisoformat(report_date) if report_date else now().date()
        if day>now().date():raise ValueError("不生成未来日期日报")
        date=day.isoformat()
        with self.lock("daily"):
            previous=self.db.execute("SELECT * FROM daily_runs WHERE date=?",(date,)).fetchone()
            if previous and previous["status"]=="ok":
                path=Path(previous["report"])
                payload=json.loads(previous["payload"])
                if not path.exists():atomic_write(path,render_digest(payload))
                return {"cached":True,"report":str(path),"count":len(payload["items"]),"status":"ok"}
            errors=[]; coverage=[]; candidates={}; requests=0
            for days in (self.config["recent_days"],self.config["backfill_days"]):
                start=(day-dt.timedelta(days=days-1)).isoformat()
                for topic, cfg in self.config["topics"].items():
                    for query in cfg["queries"]:
                        requests+=1
                        try:
                            result=self.api({"search":query,"filter":f"from_publication_date:{start},to_publication_date:{date},is_retracted:false", "sort":"publication_date:desc", "per_page":50})
                            coverage.append({"topic":topic,"query":query,"from":start,"to":date,"total":result.get("meta",{}).get("count"),"retrieved":len(result.get("results",[]))})
                            for w in result.get("results",[]):
                                if not w.get("id") or w.get("is_retracted") or w.get("type") in ("book-chapter","book","paratext","editorial") or not (start <= (w.get("publication_date") or "") <= date):continue
                                w["abstract"]=abstract_text(w.get("abstract_inverted_index"))
                                score, reasons=topic_score(topic, w)
                                if score==0:continue
                                doi=(w.get("doi") or "").lower().replace("http://doi.org/","https://doi.org/")
                                canonical=doi or "title:"+norm_journal(w.get("title") or "")
                                existing=self.db.execute("SELECT id,first_seen FROM works WHERE canonical=?",(canonical,)).fetchone()
                                wid=existing["id"] if existing else w["id"]
                                first=existing["first_seen"] if existing else now().isoformat()
                                with self.db:self.db.execute("INSERT INTO works VALUES(?,?,?,?,?,?) ON CONFLICT(id) DO UPDATE SET title=excluded.title,publication_date=excluded.publication_date,metadata=excluded.metadata",(wid,canonical,w.get("title"),w["publication_date"],first,dump(w)))
                                if self.db.execute("SELECT 1 FROM featured WHERE work_id=? AND date<>?",(wid,date)).fetchone():continue
                                journal=self.journal_match(w)
                                age=(day-dt.date.fromisoformat(w["publication_date"])).days
                                rank=score+max(0,6-age/60)+(2 if journal else 0)
                                item={"id":wid,"title":w.get("title"),"date":w["publication_date"],"type":w.get("type"),"authors":[(a.get("author") or {}).get("display_name") for a in w.get("authorships",[])[:6]],"topic":topic,"topic_name":cfg["name"],"score":round(rank,2),"reasons":reasons,"abstract":w["abstract"],"doi":doi,"url":doi or w["id"],"openalex_url":w["id"],"journal":((w.get("primary_location") or {}).get("source") or {}).get("display_name"),"journal_match":journal,"first_seen":first,"age_days":age,"basis":"OpenAlex 摘要与元数据，未核验全文" if w["abstract"] else "仅 OpenAlex 标题与元数据，缺摘要", "oa_url":(w.get("open_access") or {}).get("oa_url")}
                                if wid not in candidates or candidates[wid]["score"]<rank:candidates[wid]=item
                        except Exception as e:
                            errors.append({"topic":topic,"query":query,"error":str(e) if isinstance(e,RuntimeError) else type(e).__name__})
                            # Authentication/quota errors should not create 24 futile calls.
                            if "HTTP 401" in str(e) or "HTTP 403" in str(e) or "HTTP 429" in str(e):break
                by_topic=collections.Counter(w["topic"] for w in candidates.values())
                if all(by_topic[t]>=self.config["daily_limit"]//2 for t in self.config["topics"]) or errors:break
            chosen=[]; remaining=sorted(candidates.values(),key=lambda x:(x["score"],x["date"]),reverse=True)
            # Round-robin keeps the two interests represented, without duplicating the same DOI.
            while remaining and len(chosen)<self.config["daily_limit"]:
                for topic in self.config["topics"]:
                    item=next((x for x in remaining if x["topic"]==topic),None)
                    if item and len(chosen)<self.config["daily_limit"]:chosen.append(item);remaining.remove(item)
            report=self.root/"reports"/str(day.year)/(date+".md")
            payload={"date":date,"timezone":"Asia/Shanghai","generated_at":now().isoformat(),"items":chosen,"coverage":coverage,"errors":errors,"request_count":requests,"candidate_count":len(candidates),"status":"partial" if errors else "ok"}
            atomic_write(report,render_digest(payload))
            atomic_write(report.with_suffix(".json"),dump(payload))
            with self.db:
                self.db.execute("DELETE FROM featured WHERE date=?",(date,))
                self.db.executemany("INSERT INTO featured VALUES (?,?,?)",[(x["id"],date,x["topic"]) for x in chosen])
                self.db.execute("INSERT OR REPLACE INTO daily_runs VALUES (?,?,?,?)",(date,payload["status"],str(report),dump(payload)))
            return {"cached":False,"report":str(report),"json":str(report.with_suffix(".json")),"count":len(chosen),"status":payload["status"],"errors":errors,"request_count":requests}


def split_chunks(text, target=850):
    sentences=re.split(r"(?<=[.!?。！？])\s+|(?<=[。！？])(?=[\u4e00-\u9fff])",text)
    current=""
    for s in sentences:
        if current and len(current)+len(s)>target:
            if len(current)>35:yield current
            current=""
        if len(s)>target*2:
            if current:yield current;current=""
            for i in range(0,len(s),target):
                if len(s[i:i+target])>35:yield s[i:i+target]
        else:current=(current+" "+s).strip()
    if len(current)>35:yield current


def query_terms(query):
    groups=[];terms=[]
    for zh, vals in CONCEPTS.items():
        if zh in query or any(v.lower() in query.lower() for v in vals if v.isascii()):groups.append(vals);terms.extend(vals)
    for word in re.findall(r"[A-Za-z][A-Za-z0-9_-]{2,}",query):
        if word.lower() not in {"the","and","with","for","advantages","limitations","methods","results","what","how"}:terms.append(word)
    if not terms:
        for seq in re.findall(r"[\u4e00-\u9fff]+",query):
            for stop in ["有什么","技术路线","优缺点","缺点","优势","帮我","找一下","局限","的","请"]:seq=seq.replace(stop,"")
            terms.extend(seq[i:i+2] for i in range(max(0,len(seq)-1)))
    # Chinese terms longer than two chars are indexed as bigrams for recall.
    indexed=[]
    for term in terms:
        if re.fullmatch(r"[\u4e00-\u9fff]{3,}",term):indexed.extend(term[i:i+2] for i in range(len(term)-1))
        else:indexed.append(term)
    return list(dict.fromkeys(indexed))[:45],groups


def best_sentence(text,role,terms):
    ss=re.split(r"(?<=[.!?。！？])\s+",text)
    cues=ROLE_CUES.get(role,[])
    choices=[s for s in ss if len(s)>40]
    if not choices:return text[:500]
    return max(choices,key=lambda s:5*sum(c in s.lower() for c in cues)+sum(t.lower() in s.lower() for t in terms))


def abstract_text(inv):
    if not inv:return ""
    ordered={p:word for word,positions in inv.items() for p in positions}
    return " ".join(ordered[k] for k in sorted(ordered))


def topic_score(topic,w):
    title=normalize(w.get("title") or "").lower();text=title+" "+normalize(w.get("abstract","")).lower()
    core_boost=0
    if topic=="gas":
        # Thermal insulating gas in oil wells/buildings is not electrical insulation.
        if not any(x in text for x in ["dielectric","electrical","breakdown","discharge","high-voltage","high voltage","sf6","sf 6","sulfur hexafluoride","绝缘","放电"]):return 0,[]
        domain=[x for x in ["insulating gas","insulation gas","dielectric gas","sf6 alternative","sf6 replacement","fluoronitrile","fluoroketone","c4f7n","c5f10o","c6f12o","绝缘气体","环保气体"] if x in text]
        if not domain and (("sf6" in text or "sulfur hexafluoride" in text) and any(x in text for x in ["breakdown","dielectric strength","insulation"])):domain=["SF6 + 电气绝缘"]
        if not domain:return 0,[]
        method=[x for x in ["molecular","descriptor","machine learning","structure-property","screening","design","quantum","dielectric strength","breakdown","electron attach","分子","构效"] if x in text]
        if not method:return 0,[]
        molecular=[x for x in ["molecular structure","structure-property","descriptor","molecular design","density functional","dft","first-principle","quantum","qspr","构效"] if x in text]
        if molecular:core_boost=12;method=molecular+method
        if any(x in title for x in ["sensing","sensor","nanomaterial"]):core_boost-=7
    else:
        domain=[x for x in ["partial discharge","局部放电","gas-insulated","high-voltage","high voltage"] if x in text]
        if not domain and "transformer" in text and any(x in text for x in ["insulation","fault diagnosis"]):domain=["transformer fault/insulation"]
        if not domain:return 0,[]
        method=[x for x in ["fusion","multimodal","multi-modal","multi-sensor","acoustic","optical","ultrasonic","prognostic","prediction","pattern recognition","diagnosis","severity","classification","识别","融合"] if x in text]
        if not method:return 0,[]
    title_boost=3 if any(x in title for x in domain) else 0
    return 4+min(len(method),6)+title_boost+core_boost,domain[:3]+method[:5]


def render_digest(payload):
    day=payload["date"];items=payload["items"]
    lines=[f"# 高电压研究文献日报｜{day}","",f"> 北京时间｜入选 {len(items)} 篇｜数据来源：OpenAlex｜生成：{payload['generated_at']}","", "本报告区分近 7 天发表与较早补充阅读；首次发现不等于当天发表。期刊层级来自用户提供的学院目录，不代表当前 JCR/中科院分区。",""]
    if not items:lines += ["今日未筛到可用的未推送条目。覆盖有限，不能据此断言没有新论文。",""]
    for topic in ["gas","fault"]:
        group=[x for x in items if x["topic"]==topic]
        if not group:continue
        lines += ["## "+group[0]["topic_name"],""]
        for i,item in enumerate(group,1):
            fresh="近 7 天发表" if item["age_days"]<=6 else "补充阅读（较早发表，本次首次推荐）"
            lines += [f"### {i}. [{item['title']}]({item['url']})", "",f"- **发表日期：** {item['date']}；{fresh}。",f"- **期刊：** {item['journal'] or '未收录'}。"]
            j=item.get("journal_match")
            if j:lines += [f"- **学院目录：** {j['tier']}；{j['source_sheet']}!B{j['source_row']}:G{j['source_row']}（按 ISSN 或刊名匹配）。"]
            lines += [f"- **关注原因：** 检索匹配 {', '.join(item['reasons'])}。",f"- **证据范围：** {item['basis']}。"]
            if item["abstract"]:lines += [f"- **摘要节选（原文）：** {item['abstract'][:650]}{'…' if len(item['abstract'])>650 else ''}"]
            else:lines += ["- **摘要：** 未提供，不能据标题推断实验效果。"]
            if item.get("oa_url"):lines += [f"- **开放获取入口：** [查看全文入口]({item['oa_url']})（尚未下载入本地库）。"]
            lines += [f"- **索引：** [OpenAlex 原始记录]({item['openalex_url']})。",""]
    lines += ["## 覆盖与运行状态","",f"本轮 {payload['request_count']} 次检索，去重后候选 {payload['candidate_count']} 篇；每条查询最多读取 50 篇。默认检索近 30 天，不足时扩展到 365 天。按 DOI/OpenAlex ID 去重，并排除已推荐条目、撤稿和晚于报告日期的记录。", "", "OpenAlex 可能延迟收录或缺少摘要；未覆盖全部出版社与中文数据库。本地 PDF 全文和远程题录是分开存储的。"]
    if payload["errors"]:lines += ["", "本次为部分成功："]+[f"- {x['topic']} / {x['query']}：{x['error']}" for x in payload["errors"]]
    return "\n".join(lines)+"\n"


def write_evidence(lib,query,role,limit=8):
    result=lib.search(query,role,limit)
    lines=["# 写作证据检索", "",f"需求：{query}；段落意图：{ROLE_ZH[role]}。", "", "以下为真实 PDF 文字层摘录。句式组织属于写作建议，须与原文结论、条件及适用范围分开。", ""]
    if not result["results"]:lines += ["没有找到满足条件的本地证据。请缩小关键词或补充该主题 PDF；不能用通用模板代替文献结论。",""]
    for i,d in enumerate(result["results"],1):
        lines += [f"## {i}. {d['title']}","",f"位置：PDF 第 {d['page']} 页；证据编号 `{d['document_id'][:12]}:p{d['page']}:c{d['id']}`。", "",f"[打开原文件](<{d['paths'][0]}>)", "",f"> {d['quote']}", "", "上下文：", "",d["text"],""]
    lines += ["## 组织段落时的建议", "", "可按“技术适用的任务 → 已证实的优势 → 成立条件 → 具体局限 → 本文拟解决的问题”组织。该结构是写作建议，不是上述任一论文的原句。引用时保留每个结论的来源，不把不同论文的适用条件拼接成新的事实。", ""]
    path=lib.root/"outputs"/("写作证据_"+now().strftime("%Y%m%d_%H%M%S")+".md")
    atomic_write(path,"\n".join(lines))
    return {"path":str(path),"hits":len(result["results"]),"results":result["results"]}


TOOLS=[
 {"name":"library_status","description":"本地论文库数量、路径、解析质量与索引状态。","inputSchema":{"type":"object","properties":{}}},
 {"name":"search_passages","description":"按技术主题与段落意图检索真实 PDF 原句，返回出处、PDF页码、上下文。不是生成文本。","inputSchema":{"type":"object","properties":{"query":{"type":"string"},"role":{"type":"string","enum":list(ROLE_ZH)},"limit":{"type":"integer","minimum":1,"maximum":30}},"required":["query"]}},
 {"name":"read_page","description":"回读指定论文完整页文字，核验片段条件与上下文。page 为从1开始的PDF物理页。","inputSchema":{"type":"object","properties":{"document_id":{"type":"string"},"page":{"type":"integer","minimum":1}},"required":["document_id","page"]}},
 {"name":"sync_library","description":"增量索引已配置论文目录；只读取原PDF，更新本地SQLite。","inputSchema":{"type":"object","properties":{}}},
 {"name":"daily_digest","description":"检索OpenAlex两个高电压主题，去重并保存本地日报；同日成功任务幂等，不发送外部消息。","inputSchema":{"type":"object","properties":{}}}
]


def mcp(lib):
    """MCP stdio JSON-RPC with no stdout logging; compatible with DSH's MCP SDK."""
    for line in sys.stdin:
        req=None
        try:
            req=json.loads(line)
            if "id" not in req:continue
            method=req.get("method");args=req.get("params",{})
            if method=="initialize":
                version=args.get("protocolVersion", "2024-11-05")
                if version not in ("2024-11-05","2025-03-26","2025-06-18"):version="2025-06-18"
                out={"protocolVersion":version,"capabilities":{"tools":{"listChanged":False}},"serverInfo":{"name":"hv-literature","version":VERSION}}
            elif method=="ping":out={}
            elif method=="tools/list":out={"tools":TOOLS}
            elif method=="tools/call":
                name=args.get("name");a=args.get("arguments",{})
                try:
                    if name=="library_status":value=lib.status()
                    elif name=="search_passages":value=lib.search(**a)
                    elif name=="read_page":value=lib.page(**a)
                    elif name=="sync_library":value=lib.ingest()
                    elif name=="daily_digest":value=lib.daily()
                    else:raise ValueError("未知工具")
                    out={"content":[{"type":"text","text":dump(value)}],"isError":False}
                except Exception as e:out={"content":[{"type":"text","text":str(e) if isinstance(e,(ValueError,RuntimeError)) else type(e).__name__}],"isError":True}
            else:
                print(json.dumps({"jsonrpc":"2.0","id":req["id"],"error":{"code":-32601,"message":"Method not found"}}),flush=True);continue
            print(json.dumps({"jsonrpc":"2.0","id":req["id"],"result":out},ensure_ascii=False),flush=True)
        except Exception:
            print(json.dumps({"jsonrpc":"2.0","id":req.get("id") if isinstance(req,dict) else None,"error":{"code":-32700,"message":"Invalid request"}}),flush=True)


def initialize(root, library_roots, journal_file=None):
    root = Path(root).expanduser().resolve()
    root.mkdir(parents=True, exist_ok=True)
    path = root / "config.json"
    if path.exists():
        raise RuntimeError("config.json 已存在；请直接编辑，初始化不会覆盖已有配置")
    config = json.loads((ROOT / "config.example.json").read_text(encoding="utf-8"))
    config["library_roots"] = [str(Path(p).expanduser().resolve()) for p in library_roots]
    if journal_file:
        config["journal_file"] = str(Path(journal_file).expanduser().resolve())
    # Exclusive creation protects an existing configuration even across concurrent runs.
    with path.open("x", encoding="utf-8") as f:
        f.write(dump(config) + "\n")
    return {"config":str(path),"library_roots":config["library_roots"],"next":"运行 ingest 索引论文；日报另需 OpenAlex 凭据"}


def main():
    p=argparse.ArgumentParser(description="高电压研究文献助手")
    p.add_argument("--root",default=str(ROOT));sub=p.add_subparsers(dest="cmd",required=True)
    s=sub.add_parser("init")
    s.add_argument("--library-root",action="append",required=True,help="本地论文目录，可重复指定")
    s.add_argument("--journal-file",help="可选：学院期刊目录 XLSX")
    for name in ["status","ingest","journals","daily","mcp","dashboard"]:sub.add_parser(name)
    for name in ["search","evidence"]:
        s=sub.add_parser(name);s.add_argument("query");s.add_argument("--role",choices=list(ROLE_ZH),default="all");s.add_argument("--limit",type=int,default=8)
    s=sub.add_parser("page");s.add_argument("document_id");s.add_argument("page",type=int)
    a=p.parse_args()
    if a.cmd=="init":
        print(dump(initialize(a.root,a.library_root,a.journal_file)));return
    lib=Library(a.root)
    if a.cmd=="mcp":mcp(lib);return
    if a.cmd=="status":out=lib.status()
    elif a.cmd=="ingest":out=lib.ingest()
    elif a.cmd=="journals":out=lib.import_journals()
    elif a.cmd=="daily":out=lib.daily()
    elif a.cmd=="search":out=lib.search(a.query,a.role,a.limit)
    elif a.cmd=="evidence":out=write_evidence(lib,a.query,a.role,a.limit)
    elif a.cmd=="page":out=lib.page(a.document_id,a.page)
    elif a.cmd=="dashboard":
        from dashboard import build_dashboard
        out=build_dashboard(lib)
    print(dump(out))


if __name__=="__main__":
    try:main()
    except Exception as e:
        print(dump({"error":str(e) if isinstance(e,(ValueError,RuntimeError)) else type(e).__name__}),file=sys.stderr)
        sys.exit(1)
